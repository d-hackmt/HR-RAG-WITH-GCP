"""Read the raw HR policy documents from Cloud Storage.

Ingests the HR policy documents from a GCS bucket, tagging each with a
`policy_category` pulled straight out of its own "Policy Category: ..."
line — that metadata is what metadata filtering (and the scope guardrail)
filter on.

For the reliability path, raw files (any format: .txt/.pdf/.docx/.pptx)
are parsed to plain text once by hr_assistant/processor.py and written out
as JSON under processed/ (see hr_assistant/ingestion.py).
`load_processed_documents_from_gcs()` below — the "JSON loader" — is what
the ingestion pipeline reads from. Parsing raw files is a one-time cost;
the JSON zone is what gets chunked and embedded.
"""

import io
import json

from docx import Document as DocxReader
from google.cloud import storage
from langchain_core.documents import Document
from pptx import Presentation
from pypdf import PdfReader

from hr_assistant import config


def extract_policy_category(text: str) -> str:
    """Every policy file starts with a 'Policy Category: X' line (HR docs)
    or a 'Category: X' line (noise docs) — pull whichever is present out."""
    for line in text.splitlines()[:5]:
        stripped = line.strip().lower()
        if stripped.startswith("policy category:"):
            return line.split(":", 1)[1].strip()
        if stripped.startswith("category:"):
            return line.split(":", 1)[1].strip()
    return "Unknown"


def _parse_pdf(raw_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(raw_bytes))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _parse_docx(raw_bytes: bytes) -> str:
    doc = DocxReader(io.BytesIO(raw_bytes))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_pptx(raw_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(raw_bytes))
    lines = []
    for slide in prs.slides:
        if slide.shapes.title is not None:
            lines.append(slide.shapes.title.text)
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                lines.append(shape.text_frame.text)
    return "\n".join(lines)


_PARSERS = {
    ".pdf": _parse_pdf,
    ".docx": _parse_docx,
    ".pptx": _parse_pptx,
}


def parse_blob(blob) -> str:
    """Download a blob and return its plain-text content, dispatching on
    file extension. .txt decodes directly; other formats go through the
    matching parser above."""
    ext = "." + blob.name.rsplit(".", 1)[-1].lower() if "." in blob.name else ""
    if ext == ".txt":
        return blob.download_as_text()
    parser = _PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"No parser registered for file extension {ext!r} ({blob.name})")
    return parser(blob.download_as_bytes())


def load_documents_from_gcs(
    bucket_name: str = config.GCS_BUCKET_NAME,
    prefix: str = config.GCS_PREFIX,
) -> list[Document]:
    """List and download every .txt file under the given GCS prefix,
    returning one LangChain Document per file with source + category metadata."""
    client = storage.Client(project=config.PROJECT_ID)
    bucket = client.bucket(bucket_name)

    documents = []
    for blob in bucket.list_blobs(prefix=prefix):
        if not blob.name.endswith(".txt"):
            continue
        text = blob.download_as_text()
        filename = blob.name.rsplit("/", 1)[-1]
        documents.append(Document(
            page_content=text,
            metadata={
                "source": filename,
                "policy_category": extract_policy_category(text),
                "gcs_path": f"gs://{bucket_name}/{blob.name}",
            },
        ))
    return documents


def iter_parsed_raw_files(bucket, prefix):
    """Yield (blob, filename, text, policy_category) for every raw file
    under the given GCS prefix — the shared parsing loop used by
    hr_assistant/processor.py's process_raw_to_json(), so "list blobs,
    parse, extract category" only lives in one place."""
    for blob in bucket.list_blobs(prefix=prefix):
        filename = blob.name.rsplit("/", 1)[-1]
        if "." not in filename:
            continue
        text = parse_blob(blob)
        yield blob, filename, text, extract_policy_category(text)


def load_processed_documents_from_gcs(
    bucket_name: str = config.GCS_BUCKET_NAME,
    prefixes: tuple[str, ...] = (config.PROCESSED_HR_PREFIX, config.PROCESSED_NOISE_PREFIX),
) -> list[Document]:
    """The JSON loader — reads the processed zone (JSON records written by
    hr_assistant/processor.py) and returns LangChain Documents. This is
    what the ingestion pipeline calls (hr_assistant/ingestion.py) — no
    parsing happens here, just reading already-parsed text out of JSON."""
    client = storage.Client(project=config.PROJECT_ID)
    bucket = client.bucket(bucket_name)

    documents = []
    for prefix in prefixes:
        for blob in bucket.list_blobs(prefix=prefix):
            if not blob.name.endswith(".json"):
                continue
            record = json.loads(blob.download_as_text())
            documents.append(Document(
                page_content=record["text"],
                metadata={
                    "source": record["source"],
                    "policy_category": record["policy_category"],
                    "gcs_path": record["raw_gcs_path"],
                },
            ))
    return documents
